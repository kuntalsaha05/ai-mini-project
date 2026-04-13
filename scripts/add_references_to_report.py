#!/usr/bin/env python3
"""
Update the report with research paper references.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
REPORT = ROOT / "AI_Mini_Project_Report.docx"
IMAGES = ROOT / "images"


def main():
    doc = Document(str(REPORT))

    # Add References section at the end
    doc.add_page_break()

    ref_heading = doc.add_heading("8. References", level=1)

    references = [
        (
            "Hart, P. E., Nilsson, N. J., & Raphael, B. (1968).",
            '"A Formal Basis for the Heuristic Determination of Minimum Cost Paths". IEEE Transactions on Systems Science and Cybernetics, 4(2), 100-107.',
            "Classic foundational work on A* algorithm."
        ),
        (
            "Dijkstra, E. W. (1959).",
            '"A note on two problems in connexion with graphs". Numerische Mathematik, 1(1), 269-271.',
            "Original Dijkstra shortest-path algorithm."
        ),
        (
            "Watkins, C. J., & Dayan, P. (1992).",
            '"Q-learning". Machine Learning, 8(3), 279-292.',
            "Foundation for Q-learning reinforcement learning used in maze solver."
        ),
        (
            "Breadth-First Search (BFS) & Depth-First Search (DFS).",
            "Cormen, T. H., Leiserson, C. E., Rivest, R. L., & Stein, C. (2009). Introduction to Algorithms (3rd ed.). MIT Press.",
            "Standard graph traversal algorithms."
        ),
        (
            "Educative Visualization of Algorithms.",
            "Shaer, O., & Hornecker, E. (2010). 'Tangible User Interfaces: Past, Present, and Future Directions'. Foundations and Trends in Human–Computer Interaction, 3(1–2), 1–137.",
            "Research on educational interactive visualization systems."
        ),
        (
            "Sudoku Solving Approaches.",
            "Yato, T., & Seta, T. (2003). 'Complexity and Completeness of Finding Another Solution and Its Application to Puzzles'. IEICE Transactions on Fundamentals of Electronics, Communications and Computer Sciences, E86-A(5), 1052-1060.",
            "Complexity analysis of Sudoku and constraint satisfaction approaches."
        ),
        (
            "Interactive Learning Applications.",
            "Hmelo-Silver, C. E., & Barrows, H. S. (2008). 'Facilitating Collaborative Knowledge Building'. Cognition and Instruction, 26(1), 48-94.",
            "Framework for interactive problem-solving learning environments."
        ),
    ]

    for i, (author, title, note) in enumerate(references, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Inches(0.12)

        run = p.add_run(f"[{i}] {author} ")
        run.bold = True

        run2 = p.add_run(title)
        run2.italic = True

        p.add_run(f"\n    Note: {note}")

    doc.save(str(REPORT))
    print(f"Updated report with references: {REPORT}")


if __name__ == "__main__":
    main()
