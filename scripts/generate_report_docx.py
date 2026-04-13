from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

ROOT = Path(__file__).resolve().parent.parent
IMAGES = ROOT / "images"
OUT = ROOT / "AI_Mini_Project_Report.docx"


def add_heading_paragraph(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Inches(0.08)


def add_image_with_caption(doc: Document, image_path: Path, caption: str, width: float = 6.2) -> None:
    if not image_path.exists():
        doc.add_paragraph(f"[Missing image: {image_path.name}]")
        return

    doc.add_picture(str(image_path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def main() -> None:
    doc = Document()

    title = doc.add_heading("AI Mini Project - Technical Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph(f"Generated on {date.today().isoformat()}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading_paragraph(doc, "1. Project Overview", level=1)
    add_body(
        doc,
        "AI Mini Project is an interactive web application for visualizing algorithmic problem-solving. "
        "It includes a Sudoku solver and a Maze solver with multiple search strategies and animated playback.",
    )

    add_heading_paragraph(doc, "2. Objectives", level=1)
    objectives = [
        "Provide a clear visual understanding of algorithm behavior.",
        "Allow users to compare solver strategies for pathfinding and puzzle solving.",
        "Keep the architecture lightweight and easy to deploy.",
    ]
    for item in objectives:
        doc.add_paragraph(item, style="List Bullet")

    add_heading_paragraph(doc, "3. Simplified System Architecture", level=1)
    add_body(
        doc,
        "The architecture is intentionally simple: a React frontend runs in the browser, "
        "while a Python FastAPI server serves static frontend files and a health endpoint.",
    )
    add_image_with_caption(
        doc,
        IMAGES / "system-architecture-simple.png",
        "Figure 1: Simplified system architecture",
    )

    add_heading_paragraph(doc, "4. Core Features", level=1)
    features = [
        "Sudoku solver with hints, scoring, and difficulty presets.",
        "Maze solver with BFS, DFS, A*, Dijkstra, Greedy, and Q-learning modes.",
        "Animation controls: run, play/pause, rewind, and fast-forward.",
        "Weighted maze mode to demonstrate cost-aware search behavior.",
    ]
    for item in features:
        doc.add_paragraph(item, style="List Bullet")

    add_heading_paragraph(doc, "5. Experimental Results (Bar Graphs)", level=1)
    add_body(
        doc,
        "The following charts summarize representative performance and behavior trends for poster/report usage.",
    )

    add_image_with_caption(
        doc,
        IMAGES / "poster-bar-maze-runtime.png",
        "Figure 2: Maze algorithm runtime comparison",
    )
    add_image_with_caption(
        doc,
        IMAGES / "poster-bar-maze-explored-cells.png",
        "Figure 3: Cells explored by maze algorithms",
    )
    add_image_with_caption(
        doc,
        IMAGES / "poster-bar-maze-path-steps.png",
        "Figure 4: Path steps by maze algorithms",
    )
    add_image_with_caption(
        doc,
        IMAGES / "poster-bar-sudoku-solve-time.png",
        "Figure 5: Sudoku solve time by difficulty",
    )

    add_heading_paragraph(doc, "6. Conclusion", level=1)
    add_body(
        doc,
        "The project demonstrates how algorithm education can be made interactive and practical using "
        "a browser-first architecture. The system remains easy to maintain while still supporting rich visual analytics.",
    )

    add_heading_paragraph(doc, "7. Future Scope", level=1)
    future = [
        "User accounts and global leaderboard.",
        "Additional maze generation strategies and benchmark automation.",
        "Mobile-first UX optimization and expanded educational walkthroughs.",
    ]
    for item in future:
        doc.add_paragraph(item, style="List Bullet")

    doc.save(OUT)
    print(f"Report created: {OUT}")


if __name__ == "__main__":
    main()
