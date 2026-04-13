from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

ROOT = Path(__file__).resolve().parent.parent
IMAGES = ROOT / "images"
OUT = ROOT / "AI_Mini_Project_Report_with_References.docx"


def add_heading_paragraph(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Inches(0.08)


def add_image_with_caption(doc: Document, image_path: Path, caption: str, width: float = 6.2) -> None:
    if not image_path.exists():
        doc.add_paragraph(f"[Missing image: {image_path.name}]")
        return

    doc.add_picture(str(image_path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def main() -> None:
    doc = Document()

    title = doc.add_heading("AI Mini Project - Technical Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph(f"Generated on {date.today().isoformat()}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading_paragraph(doc, "1. Project Overview", level=1)
    add_body(
        doc,
        "AI Mini Project is an interactive web application for visualizing algorithmic problem-solving. "
        "It includes a Sudoku solver and a Maze solver with multiple search strategies and animated playback.",
    )

    add_heading_paragraph(doc, "2. Objectives", level=1)
    objectives = [
        "Provide a clear visual understanding of algorithm behavior.",
        "Allow users to compare solver strategies for pathfinding and puzzle solving.",
        "Keep the architecture lightweight and easy to deploy.",
    ]
    for item in objectives:
        doc.add_paragraph(item, style="List Bullet")

    add_heading_paragraph(doc, "3. Simplified System Architecture", level=1)
    add_body(
        doc,
        "The architecture is intentionally simple: a React frontend runs in the browser, "
        "while a Python FastAPI server serves static frontend files and a health endpoint.",
    )
    add_image_with_caption(
        doc,
        IMAGES / "system-architecture-simple.png",
        "Figure 1: Simplified system architecture",
    )

    add_heading_paragraph(doc, "4. Core Features", level=1)
    features = [
        "Sudoku solver with hints, scoring, and difficulty presets.",
        "Maze solver with BFS, DFS, A*, Dijkstra, Greedy, and Q-learning modes.",
        "Animation controls: run, play/pause, rewind, and fast-forward.",
        "Weighted maze mode to demonstrate cost-aware search behavior.",
    ]
    for item in features:
        doc.add_paragraph(item, style="List Bullet")

    add_heading_paragraph(doc, "5. Experimental Results (Bar Graphs)", level=1)
    add_body(
        doc,
        "The following charts summarize representative performance and behavior trends for poster/report usage.",
    )

    add_image_with_caption(
        doc,
        IMAGES / "poster-bar-maze-runtime.png",
        "Figure 2: Maze algorithm runtime comparison",
    )
    add_image_with_caption(
        doc,
        IMAGES / "poster-bar-maze-explored-cells.png",
        "Figure 3: Cells explored by maze algorithms",
    )
    add_image_with_caption(
        doc,
        IMAGES / "poster-bar-maze-path-steps.png",
        "Figure 4: Path steps by maze algorithms",
    )
    add_image_with_caption(
        doc,
        IMAGES / "poster-bar-sudoku-solve-time.png",
        "Figure 5: Sudoku solve time by difficulty",
    )

    add_heading_paragraph(doc, "6. Conclusion", level=1)
    add_body(
        doc,
        "The project demonstrates how algorithm education can be made interactive and practical using "
        "a browser-first architecture. The system remains easy to maintain while still supporting rich visual analytics.",
    )

    add_heading_paragraph(doc, "7. Future Scope", level=1)
    future = [
        "User accounts and global leaderboard.",
        "Additional maze generation strategies and benchmark automation.",
        "Mobile-first UX optimization and expanded educational walkthroughs.",
    ]
    for item in future:
        doc.add_paragraph(item, style="List Bullet")

    # Add References section
    doc.add_page_break()
    add_heading_paragraph(doc, "8. References", level=1)

    references = [
        {
            "author": "Hart, P. E., Nilsson, N. J., & Raphael, B. (1968).",
            "title": '"A Formal Basis for the Heuristic Determination of Minimum Cost Paths"',
            "venue": "IEEE Transactions on Systems Science and Cybernetics, 4(2), 100-107.",
            "note": "Classic foundational work on the A* algorithm.",
        },
        {
            "author": "Dijkstra, E. W. (1959).",
            "title": '"A note on two problems in connexion with graphs"',
            "venue": "Numerische Mathematik, 1(1), 269-271.",
            "note": "Original shortest-path algorithm (Dijkstra's algorithm).",
        },
        {
            "author": "Watkins, C. J., & Dayan, P. (1992).",
            "title": '"Q-learning"',
            "venue": "Machine Learning, 8(3), 279-292.",
            "note": "Foundation for Q-learning reinforcement learning used in maze solver.",
        },
        {
            "author": "Cormen, T. H., Leiserson, C. E., Rivest, R. L., & Stein, C. (2009).",
            "title": '"Introduction to Algorithms"',
            "venue": "3rd Edition. MIT Press.",
            "note": "Standard reference covering BFS, DFS, and other graph traversal algorithms.",
        },
        {
            "author": "Shaer, O., & Hornecker, E. (2010).",
            "title": '"Tangible User Interfaces: Past, Present, and Future Directions"',
            "venue": "Foundations and Trends in Human–Computer Interaction, 3(1–2), 1–137.",
            "note": "Research on interactive and tangible visualization systems for education.",
        },
        {
            "author": "Yato, T., & Seta, T. (2003).",
            "title": '"Complexity and Completeness of Finding Another Solution and Its Application to Puzzles"',
            "venue": "IEICE Transactions on Fundamentals of Electronics, Communications and Computer Sciences, E86-A(5), 1052-1060.",
            "note": "Complexity analysis of Sudoku solving and constraint satisfaction methods.",
        },
        {
            "author": "Hmelo-Silver, C. E., & Barrows, H. S. (2008).",
            "title": '"Facilitating Collaborative Knowledge Building"',
            "venue": "Cognition and Instruction, 26(1), 48-94.",
            "note": "Framework for interactive problem-solving learning environments.",
        },
    ]

    for i, ref in enumerate(references, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Inches(0.12)

        run = p.add_run(f"[{i}] {ref['author']} ")
        run.bold = True

        run2 = p.add_run(f"{ref['title']} ")
        run2.italic = True

        p.add_run(f"{ref['venue']} ")

        p.add_run(f"\n    Note: {ref['note']}")

    doc.save(OUT)
    print(f"Report created with references: {OUT}")


if __name__ == "__main__":
    main()
